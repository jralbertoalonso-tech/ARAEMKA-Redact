"""Manejo de libros Excel modernos (.xlsx) completamente en memoria.

Se analizan todas las hojas —también las ocultas—, sus celdas con contenido,
los nombres de pestaña y los encabezados/pies de impresión. Para que los
identificadores numéricos con etiqueta se detecten bien, el texto de cada celda
incluye como contexto su encabezado de columna o la celda situada a la izquierda.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .docx_doc import CARACTER_REDACCION
from .ooxml import buscar_textos_en_paquete, limpiar_metadatos_ooxml


@dataclass
class BloqueXlsx:
    indice: int
    texto: str
    origen: str
    referencia: str
    tipo_objetivo: str
    hoja: int
    coordenada: str | None = None
    campo: str | None = None
    inicio_valor: int = 0
    fin_valor: int = 0
    valor_visible: str = ""


def _valor_visible(celda) -> str:
    valor = celda.value
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M:%S") if valor.time() != time() \
            else valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, time):
        return valor.strftime("%H:%M:%S")
    # Conserva ceros iniciales cuando la celda usa un formato simple 000000.
    if isinstance(valor, int) and re.fullmatch(r"0+", str(celda.number_format or "")):
        return str(valor).zfill(len(celda.number_format))
    return str(valor)


def _sustituir(texto: str, spans: list[tuple]) -> str:
    trozos: list[str] = []
    cursor = 0
    for span in sorted(spans, key=lambda s: (s[0], s[1])):
        inicio = max(0, int(span[0]))
        fin = min(len(texto), int(span[1]))
        if inicio < cursor:
            inicio = cursor
        if inicio >= fin:
            continue
        reemplazo = span[2] if len(span) > 2 else None
        trozos.append(texto[cursor:inicio])
        trozos.append(reemplazo if reemplazo is not None
                      else CARACTER_REDACCION * max(4, fin - inicio))
        cursor = fin
    trozos.append(texto[cursor:])
    return "".join(trozos)


class DocumentoXlsx:
    """Un .xlsx cargado en memoria, listo para analizar y redactar."""

    _CAMPOS_CABECERA = (
        "oddHeader", "evenHeader", "firstHeader",
        "oddFooter", "evenFooter", "firstFooter",
    )

    def __init__(self, contenido: bytes):
        # La copia de trabajo ya entra sin propiedades personalizadas, enlaces
        # externos ni comentarios ocultos. El archivo original del usuario no
        # se modifica: todo esto vive únicamente en RAM.
        self.contenido = limpiar_metadatos_ooxml(contenido)
        self.bloques: list[BloqueXlsx] = []
        self._extraer()

    @staticmethod
    def _abrir(contenido: bytes):
        return load_workbook(
            io.BytesIO(contenido), read_only=False, data_only=False, keep_links=False
        )

    def _extraer(self):
        libro = self._abrir(self.contenido)
        indice = 0
        for num_hoja, hoja in enumerate(libro.worksheets):
            estado = " (oculta)" if hoja.sheet_state != "visible" else ""
            self.bloques.append(BloqueXlsx(
                indice, hoja.title, "hoja", f"Pestaña {num_hoja + 1}{estado}",
                "titulo_hoja", num_hoja, inicio_valor=0,
                fin_valor=len(hoja.title), valor_visible=hoja.title,
            ))
            indice += 1

            encabezados: dict[int, str] = {}
            if hoja.max_row >= 1:
                for celda in hoja[1]:
                    visible = _valor_visible(celda).strip()
                    if visible and not visible.startswith("="):
                        encabezados[celda.column] = visible

            for fila in hoja.iter_rows():
                for celda in fila:
                    visible = _valor_visible(celda)
                    if not visible.strip():
                        continue
                    contexto = ""
                    if celda.row > 1:
                        contexto = encabezados.get(celda.column, "").strip()
                    if not contexto and celda.column > 1:
                        izquierda = hoja.cell(celda.row, celda.column - 1)
                        contexto = _valor_visible(izquierda).strip()
                    if contexto == visible.strip():
                        contexto = ""
                    prefijo = f"{contexto}: " if contexto else ""
                    texto = prefijo + visible
                    oculta = (
                        hoja.sheet_state != "visible"
                        or bool(hoja.row_dimensions[celda.row].hidden)
                        or bool(hoja.column_dimensions[get_column_letter(celda.column)].hidden)
                    )
                    origen = "formula" if celda.data_type == "f" else "celda"
                    if oculta:
                        origen += "_oculta"
                    self.bloques.append(BloqueXlsx(
                        indice, texto, origen,
                        f"{hoja.title}!{celda.coordinate}" + (" · oculta" if oculta else ""),
                        "celda", num_hoja, coordenada=celda.coordinate,
                        inicio_valor=len(prefijo), fin_valor=len(texto),
                        valor_visible=visible,
                    ))
                    indice += 1
                    if indice > 100_000:
                        raise ValueError("El libro contiene demasiadas celdas con contenido.")

            for nombre_campo in self._CAMPOS_CABECERA:
                contenedor = getattr(hoja, nombre_campo)
                for zona in ("left", "center", "right"):
                    pieza = getattr(contenedor, zona)
                    texto = pieza.text or ""
                    if not texto.strip():
                        continue
                    self.bloques.append(BloqueXlsx(
                        indice, texto,
                        "cabecera" if "Header" in nombre_campo else "pie",
                        f"{hoja.title} · {nombre_campo}.{zona}",
                        "cabecera_pie", num_hoja,
                        campo=f"{nombre_campo}.{zona}", inicio_valor=0,
                        fin_valor=len(texto), valor_visible=texto,
                    ))
                    indice += 1
        libro.close()

    def ajustar_deteccion(self, bloque: BloqueXlsx, deteccion: dict) -> dict | None:
        """Limita la detección al valor real, excluyendo el contexto añadido."""
        inicio = max(deteccion["inicio"], bloque.inicio_valor)
        fin = min(deteccion["fin"], bloque.fin_valor)
        if inicio >= fin:
            return None
        ajustada = dict(deteccion)
        ajustada["inicio"] = inicio
        ajustada["fin"] = fin
        ajustada["texto"] = bloque.texto[inicio:fin]
        return ajustada

    def spans_de_texto(self, bloque: BloqueXlsx, termino: str) -> list[tuple[int, int]]:
        if not termino.strip():
            return []
        spans = [m.span() for m in re.finditer(
            re.escape(termino.strip()), bloque.texto, re.IGNORECASE
        )]
        return [
            (max(i, bloque.inicio_valor), min(f, bloque.fin_valor))
            for i, f in spans
            if max(i, bloque.inicio_valor) < min(f, bloque.fin_valor)
        ]

    def redactar(self, spans_por_bloque: dict[int, list[tuple]]) -> bytes:
        libro = self._abrir(self.contenido)

        for indice, spans in spans_por_bloque.items():
            if not (0 <= indice < len(self.bloques)) or not spans:
                continue
            bloque = self.bloques[indice]
            relativos = []
            for span in spans:
                inicio = max(int(span[0]), bloque.inicio_valor)
                fin = min(int(span[1]), bloque.fin_valor)
                if inicio >= fin:
                    continue
                relativo = [inicio - bloque.inicio_valor, fin - bloque.inicio_valor]
                if len(span) > 2:
                    relativo.append(span[2])
                relativos.append(tuple(relativo))
            if not relativos:
                continue

            hoja = libro.worksheets[bloque.hoja]
            if bloque.tipo_objetivo == "celda":
                celda = hoja[bloque.coordenada]
                # Fechas, números y fórmulas no tienen offsets editables fiables.
                # Si se marca una parte, se sustituye el valor completo; una fecha
                # desplazada sí conserva el texto de reemplazo calculado.
                reemplazos = [s[2] for s in relativos if len(s) > 2 and s[2] is not None]
                if isinstance(celda.value, str) and celda.data_type != "f":
                    celda.value = _sustituir(bloque.valor_visible, relativos)
                elif reemplazos:
                    celda.value = reemplazos[0]
                else:
                    celda.value = CARACTER_REDACCION * max(4, len(bloque.valor_visible))
                celda.comment = None
                celda.hyperlink = None
            elif bloque.tipo_objetivo == "titulo_hoja":
                titulo = _sustituir(bloque.valor_visible, relativos)
                titulo = re.sub(r"[\\/*?:\[\]]", "-", titulo).strip()[:31] or "Hoja"
                base, intento = titulo, 2
                existentes = {h.title.casefold() for h in libro.worksheets if h is not hoja}
                while titulo.casefold() in existentes:
                    sufijo = f" ({intento})"
                    titulo = base[:31 - len(sufijo)] + sufijo
                    intento += 1
                hoja.title = titulo
            elif bloque.tipo_objetivo == "cabecera_pie":
                contenedor, zona = bloque.campo.split(".", 1)
                getattr(getattr(hoja, contenedor), zona).text = _sustituir(
                    bloque.valor_visible, relativos
                )

        # Los comentarios/notas y enlaces pueden contener información que no se
        # ve en la cuadrícula. Se eliminan en todas las hojas, hayan sido o no
        # seleccionadas celdas concretas.
        for hoja in libro.worksheets:
            for fila in hoja.iter_rows():
                for celda in fila:
                    celda.comment = None
                    celda.hyperlink = None
        if hasattr(libro, "_external_links"):
            libro._external_links = []

        propiedades = libro.properties
        for campo in (
            "creator", "lastModifiedBy", "title", "subject", "description",
            "keywords", "category", "identifier", "language", "version",
            "revision", "contentStatus",
        ):
            try:
                setattr(propiedades, campo, None)
            except (AttributeError, TypeError):
                pass

        salida = io.BytesIO()
        libro.save(salida)
        libro.close()
        return limpiar_metadatos_ooxml(salida.getvalue())

    def contiene_texto_oculto(self, textos: list[str]) -> list[str]:
        restante = "\n".join(b.texto for b in self.bloques).casefold()
        visibles = [t for t in textos if t.strip() and t.strip().casefold() in restante]
        internos = buscar_textos_en_paquete(self.contenido, textos)
        return list(dict.fromkeys(visibles + internos))
