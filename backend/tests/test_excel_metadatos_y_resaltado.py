"""Regresiones de la versión 0.10: Excel, metadatos y términos manuales."""

import io
import json
import re
import zipfile
from datetime import date
from pathlib import Path

import fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from PIL import Image

from app.almacen import ALMACEN, SesionDocumento
from app.detection.motor import MOTOR
from app.documentos.docx_doc import DocumentoDocx
from app.documentos.pdf_doc import DocumentoPdf, pdf_desde_imagen
from app.documentos.xlsx_doc import DocumentoXlsx
from app.main import app


RAIZ = Path(__file__).resolve().parents[2]
cliente = TestClient(app)


def _libro_sintetico() -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pacientes"
    hoja.append(["Paciente", "NHC", "Fecha de la toma", "Correo"])
    hoja.append(["Marta Ruiz Soler", 555123, date(2026, 9, 8), "marta@hospital.es"])
    hoja["B2"].number_format = "000000"
    hoja["A2"].comment = Comment("Revisado por Dra. Carmen Secreto", "Carmen")
    hoja["D2"].hyperlink = "mailto:marta@hospital.es"
    oculta = libro.create_sheet("Juan Perez Oculto")
    oculta.sheet_state = "hidden"
    oculta["A1"] = "Dato en hoja oculta"
    libro.properties.creator = "Dra. Carmen Secreto"
    libro.properties.lastModifiedBy = "Hospital Privado"
    libro.properties.title = "Historia de Marta Ruiz Soler"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def test_excel_extrae_contexto_y_fecha_de_toma():
    documento = DocumentoXlsx(_libro_sintetico())
    textos = [b.texto for b in documento.bloques]
    assert "Paciente: Marta Ruiz Soler" in textos
    assert "NHC: 555123" in textos
    assert "Fecha de la toma: 08/09/2026" in textos
    assert any(b.origen.endswith("oculta") and "Dato en hoja oculta" in b.texto
               for b in documento.bloques)


def test_excel_redacta_celdas_y_limpia_contenido_oculto():
    secreto = "Marta Ruiz Soler"
    documento = DocumentoXlsx(_libro_sintetico())
    spans = {}
    for termino in (secreto, "08/09/2026", "marta@hospital.es", "Juan Perez Oculto"):
        for bloque in documento.bloques:
            encontrados = documento.spans_de_texto(bloque, termino)
            if encontrados:
                spans.setdefault(bloque.indice, []).extend(encontrados)

    resultado = documento.redactar(spans)
    libro = load_workbook(io.BytesIO(resultado), data_only=False, keep_links=False)
    assert secreto not in str(libro["Pacientes"]["A2"].value)
    assert "█" in str(libro["Pacientes"]["A2"].value)
    assert "█" in str(libro["Pacientes"]["C2"].value)
    assert libro["Pacientes"]["A2"].comment is None
    assert libro["Pacientes"]["D2"].hyperlink is None
    assert "Juan Perez Oculto" not in libro.sheetnames
    libro.close()

    with zipfile.ZipFile(io.BytesIO(resultado)) as paquete:
        assert not any(n.startswith("docProps/") for n in paquete.namelist())
        assert secreto.encode() not in b"".join(paquete.read(n) for n in paquete.namelist())
    final = DocumentoXlsx(resultado)
    assert final.contiene_texto_oculto([secreto, "marta@hospital.es"]) == []


def test_excel_api_detecta_identificador_con_encabezado():
    respuesta = cliente.post(
        "/api/documentos",
        files={
            "archivo": (
                "informe.xlsx", io.BytesIO(_libro_sintetico()),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={
            "categorias": json.dumps(["nhc", "fecha", "personalizada"]),
            "lista_personalizada": json.dumps(["Marta Ruiz Soler"]),
            "lista_blanca": "[]",
        },
    )
    assert respuesta.status_code == 200, respuesta.text
    datos = respuesta.json()
    assert datos["tipo"] == "xlsx"
    assert any(d["categoria"] == "nhc" and d["texto"] == "555123"
               for d in datos["detecciones"])
    assert any(d["categoria"] == "fecha" and d["texto"] == "08/09/2026"
               for d in datos["detecciones"])
    assert any(b["referencia"] == "Pacientes!B2" for b in datos["bloques"])
    ALMACEN.eliminar(datos["id"])


def test_fecha_de_la_toma_se_detecta_y_perfil_clinico_la_activa():
    detecciones = MOTOR.detectar(
        "Fecha de la toma: 08/09/2026", categorias=["fecha"]
    )
    assert [d["texto"] for d in detecciones if d["categoria"] == "fecha"] == ["08/09/2026"]

    js = (RAIZ / "frontend/app.js").read_text(encoding="utf-8")
    perfil = re.search(r"clinico:\s*\[(.*?)\],\n\n\s*publicacion", js, re.DOTALL)
    assert perfil and '"fecha"' in perfil.group(1)


def test_busqueda_manual_devuelve_rectangulos_sobre_la_marcha():
    secreto = "Texto Que Faltaba"
    pdf = fitz.open()
    pagina = pdf.new_page()
    pagina.insert_text((72, 72), f"Antes {secreto} después")
    documento = DocumentoPdf(pdf.tobytes())
    pdf.close()
    sesion = SesionDocumento("prueba.pdf", "pdf", documento)
    ALMACEN.guardar(sesion)
    try:
        respuesta = cliente.post(
            f"/api/documentos/{sesion.id}/buscar-textos",
            json={"textos": [secreto.lower()]},
        )
        assert respuesta.status_code == 200
        coincidencias = respuesta.json()["coincidencias"]
        assert len(coincidencias) == 1
        assert coincidencias[0]["pagina"] == 0
        assert coincidencias[0]["rects"]
    finally:
        ALMACEN.eliminar(sesion.id)


def test_pdf_scrub_elimina_adjuntos_y_metadatos():
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Documento de prueba")
    pdf.set_metadata({"author": "Dra. Carmen Secreto"})
    pdf.embfile_add("historia.txt", b"Paciente Secreto", filename="historia.txt")
    documento = DocumentoPdf(pdf.tobytes())
    pdf.close()
    resultado = documento.redactar([(0, fitz.Rect(0, 0, 10, 10))])
    with fitz.open(stream=resultado, filetype="pdf") as final:
        assert final.embfile_names() == []
        assert not final.metadata.get("author")


def test_foto_se_recodifica_sin_exif_antes_de_convertirla_en_pdf():
    imagen = Image.new("RGB", (80, 60), "white")
    exif = Image.Exif()
    exif[270] = "Paciente Marta Ruiz Soler"
    exif[315] = "Dra. Carmen Secreto"
    origen = io.BytesIO()
    imagen.save(origen, format="JPEG", exif=exif)

    resultado = pdf_desde_imagen(origen.getvalue())
    assert b"Paciente Marta Ruiz Soler" not in resultado
    assert b"Dra. Carmen Secreto" not in resultado
    with fitz.open(stream=resultado, filetype="pdf") as final:
        assert len(final) == 1
        assert not final.metadata.get("author")


def test_docx_elimina_revisiones_borradas_y_propiedades():
    documento = Document()
    documento.core_properties.author = "Dra. Carmen Secreto"
    parrafo = documento.add_paragraph("Texto visible")
    borrado = OxmlElement("w:del")
    borrado.set(qn("w:author"), "Dra. Carmen Secreto")
    borrado.set(qn("w:date"), "2026-09-10T10:00:00Z")
    run = OxmlElement("w:r")
    texto = OxmlElement("w:delText")
    texto.text = "Paciente Antiguo Secreto"
    run.append(texto)
    borrado.append(run)
    parrafo._p.append(borrado)
    salida = io.BytesIO()
    documento.save(salida)

    saneado = DocumentoDocx(salida.getvalue())
    assert b"Paciente Antiguo Secreto" not in saneado.contenido
    assert b"Dra. Carmen Secreto" not in saneado.contenido
    with zipfile.ZipFile(io.BytesIO(saneado.contenido)) as paquete:
        assert not any(n.startswith("docProps/") for n in paquete.namelist())
