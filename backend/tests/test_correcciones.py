"""Pruebas de regresión de la revisión de calidad (agosto 2026).

Cada prueba corresponde a un problema real detectado en la auditoría del código
y arreglado después. Están agrupadas por área.

Ejecutar desde `backend/`:
    ../.venv/bin/python -m pytest tests/test_correcciones.py -v
"""

import io

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.detection import capa3_llm as capa3
from app.detection.motor import MOTOR, resolver_solapamientos
from app.documentos import fechas
from app.documentos.docx_doc import DocumentoDocx
from app.documentos.pdf_doc import DocumentoPdf
from app.main import app

cliente = TestClient(app)


# ══════════════════════════════════════════════════════════════════════
# Fugas de datos en la redacción
# ══════════════════════════════════════════════════════════════════════

def test_solapamiento_conserva_la_cola():
    """Si la detección ganadora queda DENTRO de otra, la cola de la perdedora
    no puede perderse: sería texto sensible sin redactar."""
    persona = "Gabriel Alfa"
    texto = f"Hospital Universitario {persona} Pruebas"
    inicio_persona = texto.index(persona)
    dets = [
        {"inicio": 0, "fin": len(texto), "texto": texto, "categoria": "centro",
         "capa": 2, "confianza": 0.5},
        {"inicio": inicio_persona, "fin": inicio_persona + len(persona),
         "texto": persona, "categoria": "persona", "capa": 1, "confianza": 0.9},
    ]
    res = resolver_solapamientos(dets, texto)
    cubierto = set()
    for d in res:
        cubierto.update(range(d["inicio"], d["fin"]))
    # Todo el intervalo original sigue cubierto por alguna detección
    faltan = [i for i in range(0, len(texto)) if i not in cubierto and not texto[i].isspace()]
    assert not faltan, f"quedó texto sin cubrir: {texto[min(faltan):max(faltan)+1]!r}"


def test_firma_sanitario_no_trunca_apellidos_con_particula():
    """«Dra. Ana de la Cruz Marrero» debe capturarse ENTERA (antes se cortaba
    en «de la» y el apellido quedaba sin redactar)."""
    dets = MOTOR.detectar("Fdo.: Dra. Ana de la Cruz Marrero", categorias=["sanitario"])
    textos = [d["texto"] for d in dets]
    assert any("Cruz Marrero" in t for t in textos), textos


def test_docx_redacta_texto_de_hipervinculos():
    """Un email dentro de un hipervínculo también debe desaparecer del archivo."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    p = doc.add_paragraph("Contacto: ")
    # Hipervínculo con el email como texto visible
    hl = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "ana.perez@hospital.es"
    run.append(t)
    hl.append(run)
    p._p.append(hl)
    buf = io.BytesIO()
    doc.save(buf)

    d = DocumentoDocx(buf.getvalue())
    bloque = next(b for b in d.bloques if "ana.perez" in b.texto)
    ini = bloque.texto.index("ana.perez")
    salida = d.redactar({bloque.indice: [(ini, ini + len("ana.perez@hospital.es"))]})

    assert b"ana.perez@hospital.es" not in salida
    final = DocumentoDocx(salida)
    assert "ana.perez" not in "\n".join(b.texto for b in final.bloques)


def test_docx_celdas_combinadas_no_se_duplican():
    """Las celdas combinadas devolvían el mismo párrafo dos veces y la doble
    redacción corrompía el texto."""
    doc = Document()
    tabla = doc.add_table(rows=2, cols=2)
    tabla.cell(0, 0).merge(tabla.cell(0, 1))
    tabla.cell(0, 0).text = "Paciente: Marta Ruiz Soler"
    buf = io.BytesIO()
    doc.save(buf)

    d = DocumentoDocx(buf.getvalue())
    coincidencias = [b for b in d.bloques if "Marta Ruiz" in b.texto]
    assert len(coincidencias) == 1, "la celda combinada aparece duplicada"


def test_pdf_escaneado_con_sello_de_texto_se_ocr():
    """Una página que es un escaneo con un sello de texto encima debe pasar por
    OCR igualmente (antes el sello hacía que se saltara → fuga silenciosa)."""
    # Página con una imagen a página completa + un poco de texto digital encima
    base = fitz.open()
    pg = base.new_page()
    pg.insert_text((60, 300), "Paciente: Rosa Marrero Diaz. DNI 12345678Z", fontsize=14)
    pix = pg.get_pixmap(dpi=150)
    base.close()

    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(page.rect, pixmap=pix)          # el escaneo
    page.insert_text((40, 820), "Documento validado electronicamente - Sello digital", fontsize=8)
    contenido = doc.tobytes()
    doc.close()

    d = DocumentoPdf(contenido)
    assert d.paginas[0].area_imagen_rel > 0.5
    assert 0 in d.paginas_sin_texto(), "la página escaneada con sello debe ir a OCR"


# ══════════════════════════════════════════════════════════════════════
# Capa 3 (LLM)
# ══════════════════════════════════════════════════════════════════════

def test_extraer_json_con_bloque_thinking():
    """Los modelos «thinking» (qwen3) emiten <think>…</think> antes del JSON."""
    resp = ('<think>El usuario quiere un JSON con {claves} y esas cosas.</think>\n'
            '{"entidades": [{"texto": "Ana Perez", "tipo": "persona"}]}')
    datos = capa3._extraer_json(resp)
    assert datos.get("entidades"), datos


def test_extraer_json_con_prosa_alrededor():
    resp = ('Claro, aquí tienes:\n```json\n{"entidades": [{"texto": "Luis", "tipo": "persona"}]}\n```\n'
            'Espero que te sirva {y avísame}.')
    datos = capa3._extraer_json(resp)
    assert datos.get("entidades"), datos


def test_capa3_respeta_limites_de_palabra(monkeypatch):
    """«Ana» no debe marcar «Anamnesis» ni «Analítica»."""
    texto = "Anamnesis: sin alergias. Analitica normal. La paciente Ana acude sola."

    def fake_post(url, cuerpo, timeout):
        return {"choices": [{"message": {"content": '{"entidades":[{"texto":"Ana","tipo":"persona"}]}'}}]}

    monkeypatch.setattr(capa3, "_post_json", fake_post)
    capa3.CONFIG.actualizar(activa=True, endpoint="http://127.0.0.1:1", modelo="m")
    dets = capa3.revisar_texto(texto, {"persona"})
    capa3.CONFIG.actualizar(activa=False)

    assert len(dets) == 1, [d["texto"] for d in dets]
    assert texto[dets[0]["inicio"]:dets[0]["fin"]] == "Ana"
    assert dets[0]["inicio"] > texto.index("paciente")


# ══════════════════════════════════════════════════════════════════════
# Falsos positivos y rangos
# ══════════════════════════════════════════════════════════════════════

def test_duracion_clinica_no_es_edad():
    """«HTA de 10 años de evolución» no es la edad del paciente."""
    dets = MOTOR.detectar(
        "Varon con HTA de 10 años de evolucion en tratamiento.",
        categorias=["fecha_nacimiento"])
    assert not [d for d in dets if "10 años" in d["texto"]], [d["texto"] for d in dets]
    # …pero la edad real sí se detecta
    dets2 = MOTOR.detectar("Paciente de 62 años que acude por dolor.",
                           categorias=["fecha_nacimiento"])
    assert any("62 años" in d["texto"] for d in dets2)


def test_rango_etario_meses_mayores_de_un_anio():
    assert fechas.rango_etario("8 meses de edad") == "menor de 1 año de edad"
    assert fechas.rango_etario("18 meses de edad") == "0-4 años de edad"


# ══════════════════════════════════════════════════════════════════════
# Seguridad y robustez de la API
# ══════════════════════════════════════════════════════════════════════

def test_capa3_rechaza_endpoints_de_internet():
    """El texto clínico no puede enviarse fuera de la red local."""
    r = cliente.post("/api/capa3/configurar",
                     json={"activa": True, "endpoint": "http://evil.example.com:9999", "modelo": "x"})
    assert r.status_code == 400
    r2 = cliente.post("/api/capa3/probar",
                      json={"endpoint": "https://api.openai.com", "modelo": "x"})
    assert r2.json()["ok"] is False
    # Una IP privada sí se admite
    r3 = cliente.post("/api/capa3/configurar",
                      json={"activa": False, "endpoint": "http://192.168.1.50:11434", "modelo": "x"})
    assert r3.status_code == 200


def test_parametros_json_invalidos_dan_error_claro():
    pdf = fitz.open()
    pdf.new_page().insert_text((70, 70), "Paciente: Pedro Armas Gonzalez")
    contenido = pdf.tobytes()
    pdf.close()
    r = cliente.post(
        "/api/documentos",
        files={"archivo": ("x.pdf", io.BytesIO(contenido), "application/pdf")},
        data={"categorias": "no-es-json", "lista_personalizada": "[]", "lista_blanca": "[]"},
    )
    assert r.status_code == 422
    assert "formato incorrecto" in r.json()["detail"].lower()
