"""Genera la FICHA TÉCNICA de ARAEMKA Redact: un resumen divulgativo (2 págs) para
presentar a asesores, colaboradores o cualquier interlocutor no técnico.
Salida: registro/FICHA_TECNICA_ARAEMKA_Redact.docx
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "registro" / "FICHA_TECNICA_ARAEMKA_Redact.docx"
AZUL = RGBColor(0x1D, 0x5D, 0x8F)


def main():
    doc = Document()
    doc.core_properties.author = "José Ramón Alberto Alonso"
    doc.core_properties.last_modified_by = "José Ramón Alberto Alonso"
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    def titulo(txt):
        p = doc.add_paragraph()
        r = p.add_run(txt)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = AZUL
        p.space_before = Pt(8)
        p.space_after = Pt(2)
        return p

    def punto(txt, negrita_hasta=None):
        p = doc.add_paragraph(style="List Bullet")
        if negrita_hasta:
            r = p.add_run(negrita_hasta)
            r.bold = True
            p.add_run(txt)
        else:
            p.add_run(txt)
        p.paragraph_format.space_after = Pt(2)
        return p

    # ── Cabecera ──────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("ARAEMKA Redact")
    r.bold = True; r.font.size = Pt(22); r.font.color.rgb = AZUL
    s = doc.add_paragraph("Ficha técnica · Herramienta de anonimización local de documentos clínicos")
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.runs[0].italic = True
    a = doc.add_paragraph("Autor: José Ramón Alberto Alonso")
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a.runs[0].font.size = Pt(9.5)

    # ── Resumen ejecutivo ─────────────────────────────────────────────────
    titulo("Resumen ejecutivo")
    doc.add_paragraph(
        "ARAEMKA Redact es una aplicación que elimina de forma automática y verificada los datos "
        "personales de informes clínicos (PDF, imágenes escaneadas y Word), para poder usarlos "
        "en investigación, docencia y publicaciones cumpliendo el RGPD. Su principio innegociable "
        "es que todo el procesamiento ocurre en local: ningún dato del paciente sale del equipo o "
        "de la red, sin conexión a internet ni servicios externos. Está desarrollada, desplegada "
        "y en funcionamiento.")

    # ── El problema ───────────────────────────────────────────────────────
    titulo("El problema que resuelve")
    doc.add_paragraph(
        "Anonimizar un informe a mano es lento y propenso a errores: es fácil olvidar un nombre, "
        "un número de historia o una fecha, y tapar con un rectángulo negro no elimina el texto "
        "subyacente (se recupera copiando y pegando). Las herramientas comerciales suelen ser de "
        "pago, en inglés y —lo más grave para datos sanitarios— envían la información a la nube.")

    # ── La solución ───────────────────────────────────────────────────────
    titulo("Qué hace")
    punto("los nombres de pacientes, familiares y personal sanitario; DNI/NIE, nº de la Seguridad "
          "Social, CIP y tarjeta sanitaria (incluido el formato T.I.S. y los patrones de Canarias); "
          "nº de historia clínica, teléfonos, direcciones, fechas, centros y servicios.",
          negrita_hasta="Detecta ")
    punto("real: elimina el texto de la capa del documento o los píxeles de la imagen, y limpia "
          "los metadatos. El dato original no se puede recuperar.", negrita_hasta="Redacción destructiva ")
    punto("por el usuario: nada se anonimiza a ciegas; se revisan las detecciones resaltadas por "
          "colores y se aceptan o descartan, con una segunda comprobación del resultado.",
          negrita_hasta="Verificación ")
    punto("de fechas para preservar la cronología clínica, sustitución de la edad exacta por rangos "
          "etarios, procesamiento por lotes e informe de auditoría por documento.",
          negrita_hasta="Opciones avanzadas: desplazamiento ")

    # ── Diferenciación ────────────────────────────────────────────────────
    titulo("Qué lo diferencia")
    for t in [
        "100 % local y sin conexión: cumple el principio de minimización del RGPD por diseño.",
        "Adaptado al español y a Canarias (identificadores, siglas hospitalarias, formatos propios).",
        "Multiformato, incluidos escaneados e imágenes mediante OCR local.",
        "Interfaz en español pensada para personal sanitario, no informático.",
        "Sin coste de licencia por uso y sin dependencia de ningún proveedor externo.",
    ]:
        punto(t)

    # ── Madurez ───────────────────────────────────────────────────────────
    titulo("Estado de madurez y validación")
    for t in [
        "Aplicación completa y funcionando, no un prototipo.",
        "Desplegada y probada sobre un servidor NAS y como ejecutable portable.",
        "Batería de 40 pruebas automáticas superadas (detección, redacción irreversible, OCR y auditoría).",
        "Métricas sobre corpus sintético: 99,6 % de detección (sensibilidad) y 98,4 % de precisión.",
        "Motor de detección en tres capas, la tercera (IA) opcional y también local.",
    ]:
        punto(t)

    # ── Despliegue ────────────────────────────────────────────────────────
    titulo("Modalidades de uso")
    punto("desde el navegador para todos los equipos de una red, sin instalar nada en ellos "
          "(ideal para PC corporativos bloqueados).", negrita_hasta="Servidor (NAS / Docker): ")
    punto("para macOS y Windows, sin instalación ni permisos de administrador, para uso fuera "
          "de la red del servidor.", negrita_hasta="Portable: ejecutable ")

    # ── Tecnología ────────────────────────────────────────────────────────
    titulo("Tecnología y licencias")
    doc.add_paragraph(
        "Construida con software libre de amplia difusión (Python, FastAPI, Microsoft Presidio, "
        "spaCy, Tesseract). ARAEMKA Redact se publica como software libre bajo GNU AGPL v3 "
        "exclusivamente, de forma compatible con PyMuPDF. La licencia permite el uso y la "
        "distribución, también de pago, y exige facilitar el código fuente correspondiente.")

    # ── Propiedad intelectual ─────────────────────────────────────────────
    titulo("Propiedad intelectual")
    doc.add_paragraph(
        "El código es original del autor y se ha preparado su inscripción en el Registro de la "
        "Propiedad Intelectual. La titularidad y cualquier vía de aprovechamiento deben valorarse "
        "según las circunstancias concretas de creación y, cuando proceda, con asesoría jurídica.")

    # ── Vías ──────────────────────────────────────────────────────────────
    titulo("Posibles vías de aprovechamiento")
    for t in [
        "Herramienta de apoyo a la investigación, la docencia y la preparación de documentos.",
        "Publicación como software abierto para la comunidad sanitaria.",
        "Servicios, soporte o distribuciones adicionales compatibles con la licencia AGPL.",
    ]:
        punto(t)

    # ── Contacto ──────────────────────────────────────────────────────────
    titulo("Contacto")
    c = doc.add_paragraph("José Ramón Alberto Alonso — ______________________________")
    c.runs[0].font.size = Pt(9.5)

    doc.save(str(SALIDA))
    print("Ficha técnica generada en:", SALIDA)


if __name__ == "__main__":
    main()
